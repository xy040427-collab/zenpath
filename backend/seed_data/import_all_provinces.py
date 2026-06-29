"""
全国 31 省数据通用导入脚本

读取各省 {省份}.xlsx，自适应列映射，导入 score_rank_table + admission_lines

用法:
  python seed_data/import_all_provinces.py              # 全量导入
  python seed_data/import_all_provinces.py --province 山东 --dry-run  # 测试单省
"""

import sys, os, asyncio, argparse, logging
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
from sqlalchemy import select, text
from app.core.database import async_session_factory
from app.models.models import ScoreRankTable, AdmissionLine

# ── 学校名清理 ──
import re

def extract_school_from_group(grp_name: str) -> str:
    """从 '北京大学第04组' 提取 '北京大学'"""
    if not grp_name:
        return ""
    # 去掉 "第XX组" 后缀
    name = re.sub(r'第\d+组.*$', '', grp_name).strip()
    # 去掉专业组后缀如 "02专业组"
    name = re.sub(r'\d+专业组$', '', name).strip()
    return name


def extract_school_from_full(full_name: str) -> str:
    """从 '北京第二外国语学院中瑞酒店管理学院(北京市,14201,独立学院)' 提取学校名"""
    if not full_name:
        return ""
    # 去掉括号内容
    name = re.sub(r'\([^)]*\)', '', full_name).strip()
    return name

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("import_all")

DATA_ROOT = str(Path(__file__).resolve().parent.parent.parent / "data" / "provinces")

# ── 省份名 → 文件名映射 ──
PROVINCE_FILE_MAP = {
    "山东": "山东", "北京": "北京", "天津": "天津", "上海": "上海",
    "浙江": "浙江", "海南": "海南",
    "河北": "河北", "辽宁": "辽宁", "江苏": "江苏", "福建": "福建",
    "湖北": "湖北", "湖南": "湖南", "广东": "广东", "重庆": "重庆",
    "黑龙江": "黑龙江", "吉林": "吉林", "安徽": "安徽", "江西": "江西",
    "广西": "广西", "贵州": "贵州", "甘肃": "甘肃",
    "山西": "山西", "河南": "河南", "陕西": "陕西", "内蒙古": "内蒙古",
    "四川": "四川", "云南": "云南", "宁夏": "宁夏", "青海": "青海",
    "新疆": "新疆", "西藏": "西藏",
}

# ── score_rank_table 列映射 ──
# 针对各省 xlsx 的不同 header，映射到标准字段
SEGMENT_COLUMN_ALIASES = {
    "year": ["year"],
    "province": ["province"],
    "category": ["category"],
    "score": ["score"],
    "segment_count": ["segment_count", "undergrad_segment_count", "本科段人数"],
    "cumulative_count": ["cumulative_count", "undergrad_cumulative_count", "本科累计人数"],
    "rank_start": ["rank_start"],
    "rank_end": ["rank_end"],
    "subject_filter": ["subject_filter", "subject_type"],
    "source_file_name": ["source_file_name"],
    "source_page_url": ["source_page_url"],
    "score_text": ["score_text"],
    "rank_type": ["rank_type"],
    "score_label": ["score_label"],
    "score_is_top_band": ["score_is_top_band"],
}


def find_column(headers: list[str], target: str) -> int | None:
    """在 headers 中查找目标列，支持别名"""
    aliases = SEGMENT_COLUMN_ALIASES.get(target, [target])
    for alias in aliases:
        for i, h in enumerate(headers):
            if h and h.strip() == alias:
                return i
    return None


def safe_int(val) -> int | None:
    if val is None: return None
    try: return int(float(str(val)))
    except: return None


def safe_str(val) -> str | None:
    if val is None: return None
    s = str(val).strip()
    return s if s else None


async def import_score_segments(province: str, filepath: Path, dry_run: bool = False):
    """导入一分一段表"""
    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    if "score_segments" not in wb.sheetnames:
        logger.warning(f"  [skip] {province} 无 score_segments sheet")
        wb.close()
        return 0

    ws = wb["score_segments"]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        wb.close()
        return 0

    headers = [str(h).strip() if h else "" for h in header_row]
    logger.info(f"  score_segments headers: {headers[:10]}")

    # 定位关键列
    col_year = find_column(headers, "year")
    col_cat = find_column(headers, "category")
    col_score = find_column(headers, "score")
    col_seg = find_column(headers, "segment_count")
    col_cum = find_column(headers, "cumulative_count")
    col_rstart = find_column(headers, "rank_start")
    col_rend = find_column(headers, "rank_end")
    col_subj = find_column(headers, "subject_filter")

    if col_year is None or col_score is None:
        logger.warning(f"  [skip] {province} 缺必要列 year/score")
        wb.close()
        return 0

    # 确定 subject_type 映射
    from app.services.province_config import get_province_config
    cfg = get_province_config(province)

    batch = []
    batch_size = 5000
    total = 0
    skipped = 0

    for row in rows_iter:
        year = safe_int(row[col_year]) if col_year < len(row) else None
        score = safe_int(row[col_score]) if col_score < len(row) else None
        if year is None or score is None:
            skipped += 1
            continue

        raw_cat = safe_str(row[col_cat]) if col_cat is not None and col_cat < len(row) else ""
        seg_count = safe_int(row[col_seg]) if col_seg is not None and col_seg < len(row) else 0
        cum_count = safe_int(row[col_cum]) if col_cum is not None and col_cum < len(row) else 0
        rstart = safe_int(row[col_rstart]) if col_rstart is not None and col_rstart < len(row) else 0
        rend = safe_int(row[col_rend]) if col_rend is not None and col_rend < len(row) else 0
        subj_filter = safe_str(row[col_subj]) if col_subj is not None and col_subj < len(row) else ""

        # 映射 subject_type
        if cfg and raw_cat:
            subject_type = cfg.segment_subject_map.get(raw_cat, raw_cat)
        elif subj_filter:
            subject_type = subj_filter
        else:
            subject_type = raw_cat or "综合"

        # rank = rank_end or cumulative_count
        rank = rend or cum_count or rstart or 0

        batch.append(ScoreRankTable(
            province=province,
            year=year,
            subject_type=subject_type,
            score=score,
            rank=rank,
            same_score_count=seg_count or 0,
            segment_count=seg_count or 0,
            cumulative_count=cum_count or 0,
            rank_start=rstart or 0,
            rank_end=rend or 0,
            data_source="official",
        ))
        total += 1

    wb.close()

    if dry_run:
        logger.info(f"  [dry-run] 将导入 {total} 行 (跳过 {skipped})")
        return total

    async with async_session_factory() as session:
        # 清空旧数据
        await session.execute(text(f"DELETE FROM score_rank_table WHERE province=:p"), {"p": province})
        await session.commit()

        for i in range(0, len(batch), batch_size):
            session.add_all(batch[i:i+batch_size])
            await session.commit()
        logger.info(f"  ✅ score_segments: {total} 行 (跳过 {skipped})")
        return total


async def import_admission_lines(province: str, filepath: Path, dry_run: bool = False):
    """导入投档线"""
    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    if "admission_lines" not in wb.sheetnames:
        logger.warning(f"  [skip] {province} 无 admission_lines sheet")
        wb.close()
        return 0

    ws = wb["admission_lines"]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        wb.close()
        return 0

    headers = [str(h).strip() if h else "" for h in header_row]
    logger.info(f"  admission_lines headers: {headers[:10]}")

    # 定位关键列（自适应）
    def _find_col(*names):
        for n in names:
            for i, h in enumerate(headers):
                if h == n:
                    return i
        return None

    col_year    = _find_col("year")
    col_cat     = _find_col("category")
    col_batch   = _find_col("batch")
    col_school  = _find_col("university_name", "college_name")
    col_major   = _find_col("major_name", "major_group_name")
    col_major_code = _find_col("major_code", "major_group_code", "college_group_code")
    col_score   = _find_col("min_score")
    col_rank    = _find_col("min_rank")
    col_plan    = _find_col("plan_count")
    col_subj    = _find_col("subject_type", "subject_filter")

    # ── 特殊列：优先用 college_group_name 提取学校名（湖北模式）──
    col_group_name = _find_col("college_group_name")
    # ── 特殊列：辽宁 college_code 作为学校名兜底 ──
    col_code_only = _find_col("college_code")

    if col_year is None:
        logger.warning(f"  [skip] {province} 缺 year 列")
        wb.close()
        return 0

    # 学校名列必须存在，或可以从 group_name 提取
    if col_school is None and col_group_name is None and col_code_only is None:
        logger.warning(f"  [skip] {province} 缺学校名/代码列 (headers: {headers[:8]})")
        wb.close()
        return 0

    from app.services.province_config import get_province_config
    cfg = get_province_config(province)

    batch = []
    batch_size = 5000
    total = 0
    skipped = 0

    for row in rows_iter:
        year = safe_int(row[col_year]) if col_year < len(row) else None
        if year is None:
            skipped += 1
            continue

        # ── 获取学校名（优先 college_group_name → college_name → college_code 兜底）──
        if col_group_name is not None:
            school = safe_str(row[col_group_name]) if col_group_name < len(row) else None
            if school:
                school = extract_school_from_group(school)
        elif col_school is not None:
            school = safe_str(row[col_school]) if col_school < len(row) else None
        else:
            school = None

        # ── 兜底：college_name 为空时用 college_code ──
        if not school and col_code_only is not None:
            raw_code = safe_str(row[col_code_only]) if col_code_only < len(row) else ""
            school = f"code_{raw_code}" if raw_code else None

        if not school:
            skipped += 1
            continue

        # ── 清理学校名 ──
        if '(' in school or '（' in school:
            school = extract_school_from_full(school)

        raw_cat = safe_str(row[col_cat]) if col_cat is not None and col_cat < len(row) else ""
        batch_name = safe_str(row[col_batch]) if col_batch is not None and col_batch < len(row) else "本科普通批"
        major_name = safe_str(row[col_major]) if col_major is not None and col_major < len(row) else ""
        major_code = safe_str(row[col_major_code]) if col_major_code is not None and col_major_code < len(row) else ""
        min_score = safe_int(row[col_score]) if col_score is not None and col_score < len(row) else 0
        min_rank = safe_int(row[col_rank]) if col_rank is not None and col_rank < len(row) else 0
        plan_count = safe_int(row[col_plan]) if col_plan is not None and col_plan < len(row) else 0
        subj_raw = safe_str(row[col_subj]) if col_subj is not None and col_subj < len(row) else ""

        # 映射 subject_type — 优先 subj_raw 当 raw_cat 为占位值时
        PLACEHOLDER_CATEGORIES = {"admission_line", "官方消息补源", "普通类本科批投档线"}
        if subj_raw and (not raw_cat or raw_cat in PLACEHOLDER_CATEGORIES):
            subject_type = subj_raw
        elif cfg and raw_cat:
            subject_type = cfg.admission_subject_map.get(raw_cat, raw_cat)
        else:
            subject_type = raw_cat or subj_raw or "综合"

        batch.append(AdmissionLine(
            year=year,
            province=province,
            subject_type=subject_type,
            university_name=school,
            major_group_code=major_code,
            major_group_name=major_name,
            batch=batch_name,
            min_score=min_score or 0,
            min_rank=min_rank or 0,
            plan_count=plan_count or 0,
            is_published=True,
        ))
        total += 1

    wb.close()

    if dry_run:
        logger.info(f"  [dry-run] 将导入 {total} 行 (跳过 {skipped})")
        return total

    async with async_session_factory() as session:
        await session.execute(text(f"DELETE FROM admission_lines WHERE province=:p"), {"p": province})
        await session.commit()

        for i in range(0, len(batch), batch_size):
            session.add_all(batch[i:i+batch_size])
            await session.commit()
        logger.info(f"  ✅ admission_lines: {total} 行 (跳过 {skipped})")
        return total


async def import_province(province: str, dry_run: bool = False):
    """导入单个省份的全部数据"""
    fname = PROVINCE_FILE_MAP.get(province, province)
    filepath = Path(DATA_ROOT) / fname / f"{fname}.xlsx"
    if not filepath.exists():
        logger.warning(f"[skip] {province}: 文件不存在 {filepath}")
        return 0, 0

    logger.info(f"\n{'='*50}\n{'[DRY-RUN] ' if dry_run else ''}导入: {province} ({filepath})\n{'='*50}")

    seg_count = await import_score_segments(province, filepath, dry_run)
    adm_count = await import_admission_lines(province, filepath, dry_run)
    return seg_count, adm_count


async def verify():
    """验证导入结果"""
    async with async_session_factory() as session:
        # 省份数
        r = await session.execute(text(
            "SELECT province, COUNT(*) as cnt FROM score_rank_table GROUP BY province ORDER BY cnt DESC"
        ))
        seg_provinces = r.fetchall()

        r = await session.execute(text(
            "SELECT province, COUNT(*) as cnt FROM admission_lines GROUP BY province ORDER BY cnt DESC"
        ))
        adm_provinces = r.fetchall()

        print("\n" + "=" * 70)
        print(f"{'省份':<10} {'一分段':>10} {'投档线':>10} {'年份范围':<15}")
        print("-" * 70)

        seg_map = {p: c for p, c in seg_provinces}
        adm_map = {p: c for p, c in adm_provinces}

        for p in sorted(set(list(seg_map.keys()) + list(adm_map.keys()))):
            s = seg_map.get(p, 0)
            a = adm_map.get(p, 0)
            # 年份范围
            yr = await session.execute(text(
                "SELECT MIN(year), MAX(year) FROM admission_lines WHERE province=:p"), {"p": p})
            yr_range = yr.fetchone()
            yr_str = f"{yr_range[0]}-{yr_range[1]}" if yr_range[0] else "—"
            print(f"{p:<10} {s:>10,} {a:>10,} {yr_str:<15}")

        total_seg = sum(seg_map.values())
        total_adm = sum(adm_map.values())
        print("-" * 70)
        print(f"{'合计':<10} {total_seg:>10,} {total_adm:>10,}")
        print(f"省份数: {len(seg_map)} (一分段), {len(adm_map)} (投档线)")


async def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--province", type=str, help="只导入指定省份")
    parser.add_argument("--dry-run", action="store_true", help="仅预览不写入")
    parser.add_argument("--verify", action="store_true", help="仅验证不导入")
    args = parser.parse_args()

    if args.verify:
        await verify()
        return

    provinces = [args.province] if args.province else list(PROVINCE_FILE_MAP.keys())

    seg_total = 0
    adm_total = 0
    failed = []

    for p in provinces:
        try:
            s, a = await import_province(p, args.dry_run)
            seg_total += s
            adm_total += a
        except Exception as e:
            logger.error(f"  ❌ {p} 导入失败: {e}")
            failed.append((p, str(e)))

    logger.info(f"\n{'='*50}")
    logger.info(f"导入完成: {len(provinces)} 省")
    logger.info(f"  一分段: {seg_total:,} 行")
    logger.info(f"  投档线: {adm_total:,} 行")
    if failed:
        logger.info(f"  失败: {len(failed)} 省")
        for p, e in failed:
            logger.info(f"    - {p}: {e}")
    logger.info(f"{'='*50}")

    if not args.dry_run:
        await verify()


if __name__ == "__main__":
    asyncio.run(main())
