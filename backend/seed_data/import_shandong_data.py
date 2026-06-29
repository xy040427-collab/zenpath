"""
山东官方数据导入脚本
从 gaokaoweb 数据仓库导入山东省 2022-2025 一分一段表和投档线数据

使用方法：python seed_data/import_shandong_data.py
"""

import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

import asyncio
import openpyxl
from sqlalchemy import select, text
from app.core.database import async_session_factory, engine, Base
from app.models.models import ScoreRankTable, AdmissionLine, SubjectRequirement

SHANDONG_XLSX = str(Path(__file__).resolve().parent.parent.parent / "data" / "provinces" / "山东" / "山东.xlsx")

# 山东选考科目→subject_type 映射（新高考3+3）
SUBJECT_TYPE_MAP = {
    '全体': '全体',
    '选考物理': 'physics',
    '选考化学': 'chemistry',
    '选考生物': 'biology',
    '选考思想政治': 'politics',
    '选考历史': 'history',
    '选考地理': 'geography',
}


async def import_score_segments():
    """导入一分一段表"""
    print("=== 导入一分一段表 ===")
    wb = openpyxl.load_workbook(SHANDONG_XLSX, data_only=True)
    ws = wb['score_segments']
    
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    total = len(rows)
    print(f"总行数: {total}")
    
    batch = []
    batch_size = 5000
    
    async with async_session_factory() as session:
        # 清空旧数据
        await session.execute(text("DELETE FROM score_rank_table WHERE province='山东'"))
        await session.commit()
        
        for i, row in enumerate(rows):
            if row[0] is None:
                continue
            # year, province, category, score, subject_filter, segment_count, cumulative_count, rank_start, rank_end, source_file, source_url
            year = int(row[0])
            province = row[1]
            category = row[2]
            score = int(row[3])
            subject_filter = row[4]
            segment_count = int(row[5]) if row[5] else 0
            cumulative_count = int(row[6]) if row[6] else 0
            rank_start = int(row[7]) if row[7] else 0
            rank_end = int(row[8]) if row[8] else 0
            
            # 位次取 rank_end（累计人数上限）
            rank = rank_end
            
            item = ScoreRankTable(
                province='山东',
                year=year,
                subject_type=subject_filter or '全体',
                score=score,
                rank=rank,
                same_score_count=segment_count,
                segment_count=segment_count,
                cumulative_count=cumulative_count,
                rank_start=rank_start,
                rank_end=rank_end,
                data_source='official',
            )
            batch.append(item)
            
            if len(batch) >= batch_size:
                session.add_all(batch)
                await session.commit()
                print(f"  已导入 {i+1}/{total} 行")
                batch = []
        
        if batch:
            session.add_all(batch)
            await session.commit()
        
        # 验证
        result = await session.execute(
            text("SELECT COUNT(*) FROM score_rank_table WHERE province='山东'")
        )
        count = result.scalar()
        print(f"✅ 一分一段表导入完成: {count} 行")


async def import_admission_lines():
    """导入投档线数据"""
    print("\n=== 导入投档线数据 ===")
    wb = openpyxl.load_workbook(SHANDONG_XLSX, data_only=True)
    ws = wb['admission_lines']
    
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    total = len(rows)
    print(f"总行数: {total}")
    
    batch = []
    batch_size = 3000
    
    async with async_session_factory() as session:
        # 清空旧数据
        await session.execute(text("DELETE FROM admission_lines WHERE province='山东'"))
        await session.commit()
        
        for i, row in enumerate(rows):
            if row[0] is None:
                continue
            # year, province, category, batch, round, major_code, major_name, major_raw, university_code, university_name, university_raw, plan_count, min_rank, min_score_estimated, rank_source_type, score_source_type, source_file, source_url
            year = int(row[0])
            province = row[1]
            category = row[2]
            batch_field = row[3] or '普通类常规批'
            round_field = row[4] or ''
            major_code = str(row[5]) if row[5] else ''
            major_name = row[6] or ''
            university_name = row[9] or ''
            plan_count = int(row[11]) if row[11] else 0
            min_rank = int(row[12]) if row[12] else 0
            min_score = int(row[13]) if row[13] else 0
            
            item = AdmissionLine(
                year=year,
                province='山东',
                subject_type='physics',  # 山东3+3不分物理/历史，通填physics
                university_name=university_name,
                major_group_code=major_code,
                major_group_name=major_name,
                batch=batch_field,
                min_score=min_score,
                min_rank=min_rank,
                plan_count=plan_count,
                is_published=True,
            )
            batch.append(item)
            
            if len(batch) >= batch_size:
                session.add_all(batch)
                await session.commit()
                print(f"  已导入 {i+1}/{total} 行")
                batch = []
        
        if batch:
            session.add_all(batch)
            await session.commit()
        
        result = await session.execute(
            text("SELECT COUNT(*) FROM admission_lines WHERE province='山东'")
        )
        count = result.scalar()
        print(f"✅ 投档线导入完成: {count} 行")


async def main():
    print("=" * 60)
    print("山东官方高考数据导入")
    print("=" * 60)
    
    try:
        await import_score_segments()
        await import_admission_lines()
        print("\n🎉 所有数据导入完成！")
    except Exception as e:
        print(f"\n❌ 导入失败: {e}")
        raise


if __name__ == "__main__":
    asyncio.run(main())
